"""docx 表格填写工具（学校/公文下发的表格类 docx 通用）。

依赖：python-docx。按本机约定用 uv 跑，不要往 Hermes venv 里装包：

    uv run --with python-docx python docx_form_fill.py dump  文件.docx        # 打印所有段落文本
    uv run --with python-docx python docx_form_fill.py dump2 文件.docx 10 30  # 打印表格第10-30行单元格
    uv run --with python-docx python docx_form_fill.py normalize 模板.docx 副本.docx

也可 import 使用：

    from docx_form_fill import (Document, normalize_syms, set_all_paras, cell_set,
                                replace_doc, dump_paragraphs, dump_cells)

写值的两条硬规则（踩过坑）：
  1) 单元格写值用**原始网格列号** table.rows[r].cells[c]（合并格自动落到同一个 _tc）；
     去掉重复 _tc 后的去重序号只能用于打印，拿来写值会静默写歪。
  2) 段落文本与键都 strip() 后比较；Word 的 run 常带行尾空格，不带 strip 会整批静默不命中。
"""

from __future__ import annotations

import sys
from copy import deepcopy

from docx import Document  # noqa: F401  (re-export for importers)
from docx.oxml.ns import qn

# Wingdings 2 码位 -> 字面字符（00A3/F0A3=box1 空框，0052=boxcheck 勾选框）
SYM2CHAR = {"00A3": "□", "F0A3": "□", "0052": "☑", "0050": "□", "009F": "□"}


# ---------------------------------------------------------------- 基础遍历

def normalize_syms(doc):
    """把所有 w:sym 符号换成字面 □ / ☑，之后才能用文本替换处理勾选。返回转换个数。"""
    n = 0
    for sym in list(doc.element.body.iter(qn("w:sym"))):
        ch = (sym.get(qn("w:char")) or "").upper()
        new = SYM2CHAR.get(ch)
        if not new:
            continue
        t = sym.makeelement(qn("w:t"), {})
        t.text = new
        sym.getparent().replace(sym, t)
        n += 1
    return n


def iter_paras(doc):
    """正文段落 + 所有表格单元格段落（合并格只出一次）。"""
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            seen = []
            for cell in row.cells:
                if any(cell._tc is s for s in seen):
                    continue
                seen.append(cell._tc)
                for p in cell.paragraphs:
                    yield p


def para_text(p):
    return "".join(r.text for r in p.runs)


# ---------------------------------------------------------------- 写文本

def ensure_run(p):
    """保证段落有 run；没有则按段落标记的字体属性新建，避免新值掉成默认字体。"""
    if p.runs:
        return
    r = p.add_run("")
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        rPr = pPr.find(qn("w:rPr"))
        if rPr is not None:
            r._r.insert(0, deepcopy(rPr))


def set_para(p, text):
    """整段重写：保留首个 run 的格式，清空其余 run。"""
    ensure_run(p)
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def set_all_paras(doc, mapping):
    """mapping: {原段落文本: 新文本}，键与段落文本均 strip() 后比较。返回命中字典。"""
    mapping = {k.strip(): v for k, v in mapping.items()}
    hit = {}
    for p in iter_paras(doc):
        key = para_text(p).strip()
        if key in mapping:
            set_para(p, mapping[key])
            hit[key] = hit.get(key, 0) + 1
    return hit


def replace_para(p, old, new):
    """跨 run 替换：先把原始偏移固化再改，避免第一个 run 变长后后续区间错位。"""
    runs = p.runs
    full = "".join(r.text for r in runs)
    idx = full.find(old)
    if idx < 0:
        return False
    spans, pos = [], 0
    for r in runs:
        spans.append((r, pos, pos + len(r.text)))
        pos += len(r.text)
    end = idx + len(old)
    first = True
    for r, s, e in spans:
        if e <= idx or s >= end:
            continue
        ls, le = max(0, idx - s), min(len(r.text), end - s)
        if first:
            r.text = r.text[:ls] + new + r.text[le:]
            first = False
        else:
            r.text = r.text[:ls] + r.text[le:]
    return True


def replace_doc(doc, old, new):
    """全文档（含表格）替换，返回命中段落数。命中 0 说明键写错或空格差异。"""
    hits = 0
    for p in iter_paras(doc):
        if old in para_text(p) and replace_para(p, old, new):
            hits += 1
    return hits


# ---------------------------------------------------------------- 表格

def row_cells(row):
    """去重后的单元格列表——**仅用于打印**，不要拿去当写值坐标。"""
    out, seen = [], []
    for c in row.cells:
        if any(c._tc is s for s in seen):
            continue
        seen.append(c._tc)
        out.append(c)
    return out


def cell_set(cell, lines):
    """按行写入单元格（保留首段首 run 格式），多余段落清空。"""
    paras = cell.paragraphs
    for i, line in enumerate(lines):
        if i < len(paras):
            set_para(paras[i], line)
        else:
            np = cell.add_paragraph(line)
            if paras:
                pPr = paras[0]._p.find(qn("w:pPr"))
                if pPr is not None:
                    np._p.insert(0, deepcopy(pPr))
            ensure_run(np)
            np.runs[0].text = line
    for j in range(len(lines), len(paras)):
        set_para(paras[j], "")


def write_cell(table, r, c, lines):
    """按原始网格列号写单元格（合并格自动落到同一个 _tc）。"""
    cell_set(table.rows[r].cells[c], lines)


# ---------------------------------------------------------------- 打印

def dump_paragraphs(doc):
    for p in doc.paragraphs:
        if para_text(p).strip():
            print(repr(para_text(p)))


def dump_cells(doc, lo=0, hi=10 ** 9):
    """打印 [T表号R行C列] '文本'，C 为原始网格列号（写值用的就是它）。"""
    for ti, t in enumerate(doc.tables):
        print("===== TABLE %d rows=%d =====" % (ti, len(t.rows)))
        for ri, row in enumerate(t.rows):
            if not (lo <= ri <= hi):
                continue
            seen = []
            for ci, cell in enumerate(row.cells):
                if any(cell._tc is s for s in seen):
                    continue
                seen.append(cell._tc)
                txt = " / ".join(x for x in (para_text(p).strip() for p in cell.paragraphs) if x)
                if txt:
                    print("R%dC%d: %r" % (ri, ci, txt))


def main(argv):
    if len(argv) < 3:
        print(__doc__)
        return 1
    cmd, path = argv[1], argv[2]
    if cmd == "dump":
        dump_paragraphs(Document(path))
    elif cmd == "dump2":
        lo = int(argv[3]) if len(argv) > 3 else 0
        hi = int(argv[4]) if len(argv) > 4 else 10 ** 9
        dump_cells(Document(path), lo, hi)
    elif cmd == "normalize":
        doc = Document(path)
        n = normalize_syms(doc)
        doc.save(argv[3])
        print("w:sym 转换 %d 个 -> %s" % (n, argv[3]))
    else:
        print(__doc__)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
